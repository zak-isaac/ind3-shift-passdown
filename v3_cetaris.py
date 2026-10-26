"""Cetaris Desktop single-file XLSX import for MST V3. No AI rewriting.

Only the selected technicians' *individual labor rows* are used; Cetaris's
Correction column deliberately is NOT read as it can combine other techs.

PM/breakdown categorization here is Primary-RFR-based (validated by this
file's own test suite, see tests/test_v3_operational_patch.py and others).
A parallel production branch independently validated SJ / PM Code as a
PM signal against a real export (10-22.xlsx, 153 rows / 9 PM rows) --
that field is now ALSO captured below (pm_code, repair_class) as an
optional, best-effort provenance field for future corroboration, but
the existing RFR-based select_records() categorization is intentionally
left unchanged pending a real side-by-side data comparison (2026-10-26
integration decision -- see integration report).
"""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

REQUIRED = (
    "WO Number", "Asset #", "WO Status", "Primary RFR",
    "Service/Repair Description", "Technician", "Technician Remarks",
)
PM_CATEGORIES = {"preventive maintenance", "preventative maintenance", "pm"}
DATE_FORMATS = ("%m/%d/%Y %I:%M %p", "%m/%d/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d")
MAX_SIZE = 12 * 1024 * 1024
MAX_ROWS = 30000


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _parse_dt(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(_text(value), fmt)
        except (TypeError, ValueError):
            pass
    return None


def parse_excel(data: bytes, stats: dict | None = None) -> list[dict]:
    """Parse a bounded, validated workbook into JSON-serializable normalized rows.

    `stats`, if given, is populated in-place with import-diagnostics counts
    (source_rows, skipped_no_wo, skipped_no_tech, parsed_rows) so an MST can
    see why a smaller-than-expected row count came through, instead of just
    inferring it. Optional and additive -- existing callers unaffected.
    """
    if stats is None:
        stats = {}
    if len(data) > MAX_SIZE:
        raise ValueError("This Excel file exceeds the 12 MB import limit.")
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Unable to read this XLSX export.") from exc
    try:
        sheet = workbook.active
        source = sheet.iter_rows(values_only=True)
        headers = next(source, None)
        if not headers:
            raise ValueError("Excel file has no headings.")
        mapping = {str(name).strip().casefold(): pos for pos, name in enumerate(headers) if name is not None}
        missing = [name for name in REQUIRED if name.casefold() not in mapping]
        if missing:
            raise ValueError("This isn't the supported Cetaris Desktop report. Missing columns: " + ", ".join(missing))
        def get(row, label):
            pos = mapping.get(label.casefold())
            return row[pos] if pos is not None and pos < len(row) else None
        results = []
        source_rows = skipped_no_wo = skipped_no_tech = 0
        for i, row in enumerate(source):
            if i >= MAX_ROWS:
                raise ValueError("Import limited to 30,000 labor rows. Narrow your Cetaris export dates.")
            if all(cell is None for cell in row):
                continue  # a genuinely blank trailing row, not real data
            source_rows += 1
            wo = _text(get(row, "WO Number"))
            tech = _text(get(row, "Technician"))
            if not wo:
                skipped_no_wo += 1
                continue
            if not tech:
                skipped_no_tech += 1
                continue
            started = _parse_dt(get(row, "Start Date"))
            # Real exports leave 'Asset ID' blank for most rows -- 'Asset #'
            # is the common case, not the fallback -- but prefer Asset ID
            # when a given export does populate it.
            asset_id = _text(get(row, "Asset ID")) or _text(get(row, "Asset #"))
            results.append({
                "wo_number": wo, "asset_id": asset_id,
                "status": _text(get(row,"WO Status")),
                "rfr": _text(get(row,"Primary RFR")),
                # Optional provenance fields, not used by select_records()'s
                # RFR-based categorization -- see module docstring.
                "pm_code": _text(get(row, "SJ / PM Code")),
                "repair_class": _text(get(row, "Repair Class")),
                "reported_issue": _text(get(row,"Service/Repair Description")),
                "tech_name": tech,
                # A WO's administrative closer can be a different person.
                # Complete By is a provenance signal, not a blanket claim.
                "completed_by": _text(get(row, "Completed By")),
                "last_technician": _text(get(row, "Last Technician")),
                "remarks": _text(get(row,"Technician Remarks")),
                "start_at": started.isoformat(timespec="seconds") if started else "",
                # 'Update Date / Time' reflects when the record was actually
                # last touched -- more reliable than WO Closed Date/Time per
                # 2026-10-23 facility feedback. Optional; absent on older
                # exports.
                "updated_at": _parse_dt(get(row, "Update Date / Time")).isoformat(timespec="seconds") if _parse_dt(get(row, "Update Date / Time")) else "",
                "opened_at": _parse_dt(get(row, "WO Date / Time")).isoformat(timespec="seconds") if _parse_dt(get(row, "WO Date / Time")) else "",
                "last_complete_at": _parse_dt(get(row, "WO Last Complete Date / Time")).isoformat(timespec="seconds") if _parse_dt(get(row, "WO Last Complete Date / Time")) else "",
                "last_shift": _text(get(row, "Last Shift")),
                "closed_at": _parse_dt(get(row, "WO Closed Date/Time")).isoformat(timespec="seconds") if _parse_dt(get(row, "WO Closed Date/Time")) else "",
                # In Service is a fallback for shift attribution, not a WO close time.
                "in_service_at": _parse_dt(get(row, "In Service Date / Time")).isoformat(timespec="seconds") if _parse_dt(get(row, "In Service Date / Time")) else "",
                "cause": _text(get(row,"Cause")),
                "work_accomplished": _text(get(row,"Work Accomplished")),
            })
        stats["source_rows"] = source_rows
        stats["skipped_no_wo"] = skipped_no_wo
        stats["skipped_no_tech"] = skipped_no_tech
        stats["parsed_rows"] = len(results)
        return results
    finally:
        workbook.close()


def technician_counts(rows: list[dict]) -> list[dict]:
    by_tech: dict[str,set] = {}
    for r in rows:
        by_tech.setdefault(r["tech_name"], set()).add(r["wo_number"])
    return [{"name":name, "count":len(wos)} for name,wos in sorted(by_tech.items())]


def select_records(rows: list[dict], selected: set[str], start: datetime, end: datetime,
                   closure_reviews: dict[str, str] | None = None):
    """Count Completed or Closed PM WOs; other Completed WOs require MST choice."""
    closure_reviews = closure_reviews or {}
    # Cetaris sometimes repeats a WO across labor rows. A confirmed Closed
    # status for that WO takes precedence over a stale Completed labor row.
    closed_wos = {r["wo_number"] for r in rows
                  if r.get("tech_name") in selected and
                  str(r.get("status","")).strip().casefold() == "closed"}
    summary = Counter()
    eligible = []
    reminder_wos = set()
    for r in rows:
        if r["tech_name"] not in selected:
            continue
        summary["selected_rows"] += 1
        # Labor Start Date is not a WO eligibility boundary.
        status = r["status"].strip().casefold()
        if status == "completed" and r["wo_number"] in closed_wos:
            status = "closed"
        category = r["rfr"].strip().casefold()
        if category == "breakdown" and status == "in process":
            summary["breakdown_in_process"] += 1
            eligible.append({**r, "kind": "oos"})
            continue
        if status == "completed" and category in {"preventive maintenance", "preventative maintenance"}:
            # Completing maintenance fulfills a PM. Cetaris administrative
            # closure may occur later and must not delay the completed count.
            summary["pm_rows"] += 1
            summary["pm_completed_not_closed"] += 1
            eligible.append({**r, "kind": "pm"})
            continue
        if status == "completed":
            # Non-PM Completed rows can be included as closure reminders
            # into passdown as a closure reminder, never as a completed PM.
            summary["completed_unclosed"] += 1
            if closure_reviews.get(r["wo_number"]) == "add":
                first_reminder = r["wo_number"] not in reminder_wos
                reminder = "Reminder: Please close this WO in Cetaris before the shift ends."
                note = r["remarks"].strip()
                if first_reminder:
                    reminder_wos.add(r["wo_number"])
                    note = (note + "; " if note else "") + reminder
                    summary["closure_reminders"] += 1
                eligible.append({**r, "kind": "happening", "closure_reminder": True,
                                 "remarks": note})
            continue
        if status != "closed":
            summary["not_closed"] += 1
            continue
        if category in {"preventive maintenance", "preventative maintenance"}:
            summary["pm_rows"] += 1
            eligible.append({**r, "kind": "pm"})
        else:
            # All other RFR categories, including unfamiliar or blank values,
            # belong in Shift Happenings. Preserve the source RFR verbatim.
            summary["happening_rows"] += 1
            eligible.append({**r, "kind": "happening"})
    # A technician may log more than one labor row on a WO. Preserve each unique remark.
    grouped = {}
    for r in eligible:
        key = (r["wo_number"], r["tech_name"], r["kind"])
        if key not in grouped:
            grouped[key] = {**r, "remarks_list": []}
        if r["remarks"] and r["remarks"] not in grouped[key]["remarks_list"]:
            grouped[key]["remarks_list"].append(r["remarks"])
        if not grouped[key].get("cause") and r.get("cause"):
            grouped[key]["cause"] = r["cause"]
        if not grouped[key]["reported_issue"] and r["reported_issue"]:
            grouped[key]["reported_issue"] = r["reported_issue"]
    records = list(grouped.values())
    summary["unique_pm_wos"] = len({r["wo_number"] for r in records if r["kind"] == "pm"})
    summary["unique_happening_wos"] = len({r["wo_number"] for r in records if r["kind"] == "happening"})
    summary["unique_oos_wos"] = len({r["wo_number"] for r in records if r["kind"] == "oos"})
    return records, dict(summary)
