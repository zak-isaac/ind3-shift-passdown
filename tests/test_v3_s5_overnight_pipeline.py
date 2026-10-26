"""End-to-end S5 cross-midnight Cetaris import and partial duplicate recovery."""
from datetime import datetime, time, timedelta
from io import BytesIO
from pathlib import Path
import re

from fastapi import FastAPI
from fastapi.responses import RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.testclient import TestClient
from openpyxl import Workbook

import db
from v3_routes import create_v3_router

ROOT = Path(__file__).resolve().parents[1]
HEADERS = [
    "WO Number", "Asset #", "WO Status", "Primary RFR",
    "Service/Repair Description", "Technician", "Technician Remarks",
    "Start Date", "WO Date / Time", "WO Last Complete Date / Time",
    "WO Closed Date/Time", "In Service Date / Time", "Last Shift", "Completed By"
]

def stamp(s):
    return datetime.fromisoformat(s) if s else None

def labor(wo, start, completed, service=None, category="Proactive Repair",
          tech="Tech S5", status="Closed", opened=None):
    return [
        wo, "M" + wo, status, category, "Maintenance work", tech,
        "Completed checks; RTS", stamp(start), stamp(opened or start),
        stamp(completed), stamp(completed), stamp(service), "S5", tech
    ]

def test_s5_saturday_shift_credits_sunday_early_work_and_imports_nonduplicates(tmp_path):
    db.set_db_path(tmp_path / "s5_overnight.db")
    db.init_db()
    shift_id = db.start_shift("2026-10-24", "S5", "Weekend MST")
    other = db.start_shift("2026-10-24", "S4", "Previous MST")
    db.upsert_pm_wo(other, "DUP-PM", "MDUP-PM", "Tech S4")

    app = FastAPI()
    app.mount("/static", StaticFiles(directory=str(ROOT / "static")), name="static")
    app.include_router(create_v3_router(
        Jinja2Templates(directory=str(ROOT / "templates")),
        is_lead_authed=lambda request: request.cookies.get("lead_auth") == "test",
        can_manage_shift=lambda request, sid: sid == shift_id,
        resolve_my_shift=lambda request: db.get_shift(shift_id),
        shift_window=lambda code, day: (
            datetime.combine(day, time(17, 30)),
            datetime.combine(day + timedelta(days=1), time(7))
        ),
        lead_redirect=lambda dest: RedirectResponse("/v3/login", status_code=303),
        now_func=lambda: datetime(2026, 10, 25, 2, 30)
    ))

    book = Workbook()
    sheet = book.active
    sheet.append(HEADERS)
    # Work from S5's evening portion.
    sheet.append(labor("S5-EVE", "2026-10-24T20:10:00",
                       "2026-10-24T20:35:00", "2026-10-24T20:30:00"))
    # Sunday 01:00 still belongs to Saturday S5.
    sheet.append(labor("S5-AM", "2026-10-25T01:10:00",
                       "2026-10-25T01:20:00", "2026-10-25T01:18:00"))
    # No labor start or In Service: WO Last Complete and matching Completed By prove finish.
    sheet.append(labor("S5-PM", None, "2026-10-25T02:00:00",
                       category="Preventive Maintenance",
                       opened="2026-10-24T12:00:00"))
    # Older WO start + verified next-day completion is attributed automatically.
    sheet.append(labor("S5-REVIEW", "2026-10-22T14:00:00",
                       "2026-10-25T02:15:00", opened="2026-10-22T13:00:00"))
    # Previously credited WO: skip it, not the rest of the file.
    sheet.append(labor("DUP-PM", "2026-10-25T01:40:00",
                       "2026-10-25T01:50:00", "2026-10-25T01:49:00",
                       category="Preventive Maintenance"))
    buffer = BytesIO()
    book.save(buffer)

    with TestClient(app) as client:
        client.cookies.set("lead_auth", "test")
        stage=client.post("/v3/import",data={"shift_id":str(shift_id)},
             files={"file":("S5_both_dates.xlsx",buffer.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert stage.status_code==200,stage.text[:700]
        assert "5 unique WOs" in stage.text
        token=re.search(r"/v3/import/([0-9a-f]{32})/confirm",stage.text).group(1)
        done=client.post(f"/v3/import/{token}/confirm",
             data={"shift_id":str(shift_id),"techs":["Tech S5"]},follow_redirects=False)
        assert done.status_code==303,done.text[:700]
        assert "Choose the owning shift" not in done.text
        screen=client.get(f"/v3?shift_id={shift_id}")
        assert screen.status_code==200
        assert "Already credited to another shift (not counted)" in screen.text
        assert "after-midnight WOs in file" in screen.text
        assert "after-midnight included" in screen.text
        assert "4</strong> after-midnight WOs in file" in screen.text
        assert "3</strong> after-midnight included" in screen.text
        assert "S5-AM" in screen.text
        assert "S5-REVIEW" in screen.text
        assert {e["wo_number"] for e in db.get_entries(shift_id)} == {
            "S5-EVE","S5-AM","S5-REVIEW"
        }
        assert db.get_pm_count(shift_id)==1
        assert db.get_pm_count(other)==1
        # Technician-supported equipment-return time is kept separately from
        # administrative closure and survives importing another XLSX later.
        assert "Oct 25, 1:18 AM" in screen.text
        assert "Next day" in screen.text
        with db._conn() as conn:
            saved=conn.execute(
                "SELECT completed_at FROM v3_wo_completion_times WHERE shift_id=? AND wo_number=?",
                (shift_id,"S5-AM")).fetchone()
        assert saved and saved[0].startswith("2026-10-25T01:18")
        next_book=Workbook()
        next_sheet=next_book.active
        next_sheet.append(HEADERS)
        next_sheet.append(labor("S5-LATER","2026-10-25T02:05:00",
                                "2026-10-25T02:20:00","2026-10-25T02:18:00"))
        next_buf=BytesIO();next_book.save(next_buf)
        staged=client.post("/v3/import",data={"shift_id":str(shift_id)},
            files={"file":("next_export.xlsx",next_buf.getvalue(),
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert staged.status_code==200
        next_token=re.search(r"/v3/import/([0-9a-f]{32})/confirm",staged.text).group(1)
        imported=client.post(f"/v3/import/{next_token}/confirm",
            data={"shift_id":str(shift_id),"techs":["Tech S5"]},follow_redirects=False)
        assert imported.status_code==303
        after=client.get(f"/v3?shift_id={shift_id}")
        assert "S5-AM" in after.text
        assert "Oct 25, 1:18 AM" in after.text
        assert "S5-LATER" in after.text



def test_s5_breakdown_in_process_is_oos_and_later_closed_becomes_rts(tmp_path,monkeypatch):
    """Full upload-to-sidebar lifecycle across Sunday -> Monday midnight."""
    monkeypatch.delenv("IND3_OLLAMA_REWRITE",raising=False)
    db.set_db_path(tmp_path/"breakdown_overnight.db"); db.init_db()
    sid=db.start_shift("2026-10-25","S5","Weekend MST")
    app=FastAPI()
    app.mount("/static",StaticFiles(directory=str(ROOT/"static")),name="static")
    app.include_router(create_v3_router(
        Jinja2Templates(directory=str(ROOT/"templates")),
        is_lead_authed=lambda req: req.cookies.get("lead_auth")=="test",
        can_manage_shift=lambda req,id:id==sid,
        resolve_my_shift=lambda req: db.get_shift(sid),
        shift_window=lambda code,day:(
            datetime.combine(day,time(17,30)),
            datetime.combine(day+timedelta(days=1),time(7))),
        lead_redirect=lambda _: RedirectResponse("/v3/login",status_code=303),
        now_func=lambda: datetime(2026,10,26,4,30)))
    def make_book(rows):
        book=Workbook(); ws=book.active;ws.append(HEADERS)
        for r in rows:ws.append(r)
        buf=BytesIO();book.save(buf)
        return buf.getvalue()
    in_process=labor("IP-400","2026-10-25T20:00:00",None,
                     status="In Process",category="Breakdown")
    # WO closing is an administrative event; actual tech return is earlier.
    first_closed=labor("CLOSED-401","2026-10-25T23:25:00",
                       "2026-10-26T01:45:00",service="2026-10-26T01:40:00",
                       status="Closed",category="Breakdown")
    with TestClient(app) as client:
        client.cookies.set("lead_auth","test")
        def import_rows(rows):
            stage=client.post("/v3/import",data={"shift_id":str(sid)},files={
                "file":("S5_Breakdown.xlsx",make_book(rows),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
            assert stage.status_code==200, stage.text[:700]
            token=re.search(r"/v3/import/([0-9a-f]{32})/confirm",stage.text).group(1)
            confirmed=client.post(f"/v3/import/{token}/confirm",
                data={"shift_id":str(sid),"techs":["Tech S5"]},follow_redirects=False)
            assert confirmed.status_code==303,confirmed.text[:700]
        import_rows([in_process,first_closed])
        assert "MIP-400" in db.get_oos_units()
        assert "MCLOSED-401" in db.get_rts_events(sid)
        display=client.get("/v3",params={"shift_id":sid})
        assert display.status_code==200
        assert "Out of Service (OOS)" in display.text
        assert "MIP-400" in display.text
        assert "MCLOSED-401" in display.text
        complete=labor("IP-400","2026-10-25T20:00:00",
                       "2026-10-26T04:30:00",service="2026-10-26T04:28:00",
                       status="Closed",category="Breakdown")
        import_rows([complete])
        assert "MIP-400" not in db.get_oos_units()
        assert set(db.get_rts_events(sid))=={"MIP-400","MCLOSED-401"}
        assert db.get_rts_events(sid).count("MIP-400")==1
        again=client.get("/v3",params={"shift_id":sid})
        assert "Oct 26, 4:28 AM" in again.text
        assert "Next day" in again.text
