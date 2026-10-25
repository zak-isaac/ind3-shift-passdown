from datetime import datetime, time, timedelta
from pathlib import Path
import re

import pytest
from fastapi import FastAPI
from fastapi.responses import RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from fastapi.testclient import TestClient

import db
from v3_routes import create_v3_router

DIR = Path(__file__).resolve().parents[1]

@pytest.fixture
def app_client(tmp_path):
    db.set_db_path(tmp_path / 'test.db')
    db.init_db()
    sid = db.start_shift('2026-10-20', 'S2', 'Test MST')
    app=FastAPI()
    app.mount('/static', StaticFiles(directory=str(DIR/'static')), name='static')
    templates=Jinja2Templates(directory=str(DIR/'templates'))
    app.include_router(create_v3_router(
        templates,
        is_lead_authed=lambda req: req.cookies.get('lead_auth')=='test',
        can_manage_shift=lambda req, candidate: candidate == sid,
        resolve_my_shift=lambda req: db.get_shift(sid),
        shift_window=lambda shift,d:(datetime.combine(d,time(0)),datetime.combine(d+timedelta(days=1),time(0))),
        lead_redirect=lambda dest: RedirectResponse('/lead/login?next='+dest,status_code=303)
    ))
    with TestClient(app) as client:
        client.cookies.set('lead_auth','test')
        yield client,sid


def test_login_and_import_and_editor_and_review(app_client):
    client,sid=app_client
    resp=client.get('/v3')
    assert resp.status_code==200, resp.text[:500]
    assert 'Shift workspace' in resp.text
    sample=Path('/mnt/data/trial.xlsx')
    if not sample.exists(): pytest.skip('trial.xlsx unavailable')
    stage=client.post('/v3/import',data={'shift_id':sid},files={'file':('trial.xlsx',sample.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert stage.status_code==200,stage.text[:500]
    assert 'Who worked this shift?' in stage.text
    import re
    token=re.search(r'/v3/import/([a-f0-9]{32})/confirm',stage.text).group(1)
    selected=['Ssaku, Isaac','Melvin, Douglas','Bawi, Mang','Popoola, Waheed']
    confirmed=client.post(f'/v3/import/{token}/confirm',data={'shift_id':str(sid),'techs':selected},follow_redirects=True)
    assert confirmed.status_code==200,confirmed.text[:700]
    assert 'VOIDFILL1201' in confirmed.text
    assert db.get_pm_count(sid)==8
    entries=db.get_entries(sid)
    assert len(entries)==4,entries
    assert {e['tech_name'] for e in entries} <=set(selected)
    e=next(e for e in entries if e['unit']=='VOIDFILL1201')
    modal=client.get(f'/v3/happening/{e["id"]}?shift_id={sid}')
    assert modal.status_code==200
    assert 'reported_issue' in modal.text and 'v3-view-fields' in modal.text
    ed=client.post(f'/v3/happening/{e["id"]}/edit',data={'shift_id':sid,'reported_issue':'Head below MIN','cause_text':'','remarks_'+str(e['id']):'Recalibrated head; RTS'},follow_redirects=False)
    assert ed.status_code==303
    assert 'Recalibrated head' in db.get_entry(e['id'])['raw_remarks']
    # Re-uploading IDENTICAL Cetaris source must not flag MST-edited text.
    stage2=client.post('/v3/import',data={'shift_id':sid},files={'file':('trial.xlsx',sample.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    token2=re.search(r'/v3/import/([a-f0-9]{32})/confirm',stage2.text).group(1)
    confirmed2=client.post(f'/v3/import/{token2}/confirm',data={'shift_id':str(sid),'techs':selected},follow_redirects=True)
    assert '0 flagged for review' in confirmed2.text
    assert 'Recalibrated head; RTS' in db.get_entry(e['id'])['raw_remarks']
    # An actual Cetaris change must create a review item without overwriting.
    from openpyxl import load_workbook
    from io import BytesIO
    w=load_workbook(BytesIO(sample.read_bytes()))
    for row in list(w.active.iter_rows(min_row=2)):
        if str(row[0].value)=='28021671':
            row[39].value='Updated technician remark from revised Cetaris export'
            break
    buf=BytesIO();w.save(buf)
    changed=client.post('/v3/import',data={'shift_id':sid},files={'file':('revised.xlsx',buf.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    token3=re.search(r'/v3/import/([a-f0-9]{32})/confirm',changed.text).group(1)
    confirmed3=client.post(f'/v3/import/{token3}/confirm',data={'shift_id':str(sid),'techs':selected},follow_redirects=True)
    assert 'Review update' in confirmed3.text
    assert 'Recalibrated head; RTS' in db.get_entry(e['id'])['raw_remarks']
    keep=client.post(f'/v3/happening/{e["id"]}/review',data={'shift_id':sid,'decision':'reject'},follow_redirects=False)
    assert keep.status_code==303
    assert 'Recalibrated head; RTS' in db.get_entry(e['id'])['raw_remarks']
    preview=client.get('/v3/preview',params={'shift_id':sid})
    assert preview.status_code==200
    assert 'Recalibrated head; RTS' in preview.text


def test_shift_ownership_and_unauthenticated_access(app_client):
    client,sid=app_client
    assert client.get('/v3',params={'shift_id':sid+999},follow_redirects=False).status_code==303
    client.cookies.clear()
    assert client.get('/v3',follow_redirects=False).status_code==303
    assert client.post('/v3/manual',data={'shift_id':sid,'unit':'M1','issue':'x'},follow_redirects=False).status_code==401


def test_changing_selected_techs_filters_pm_and_happenings(app_client):
    client,sid=app_client
    sample=Path('/mnt/data/trial.xlsx')
    if not sample.exists(): pytest.skip('trial.xlsx unavailable')
    def do_import(selected):
        staged=client.post('/v3/import',data={'shift_id':sid},files={'file':('trial.xlsx',sample.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        token=re.search(r'/v3/import/([a-f0-9]{32})/confirm',staged.text).group(1)
        return client.post(f'/v3/import/{token}/confirm',data={'shift_id':str(sid),'techs':selected},follow_redirects=True)
    first=do_import(['Ssaku, Isaac','Melvin, Douglas','Bawi, Mang','Popoola, Waheed'])
    assert db.get_pm_count(sid)==8
    second=do_import(['Ssaku, Isaac'])
    assert db.get_pm_count(sid)==2
    assert len([e for e in db.get_entries(sid) if not e['excluded']])==1
    assert 'VOIDFILL1201' in second.text


def test_reselect_tech_and_closed_pm_replacement(app_client):
    """Changing selection restores only auto-hidden entries; reopened PMs disappear."""
    import json, uuid
    client, sid = app_client
    from v3_cetaris import parse_excel
    source=Path('/mnt/data/trial.xlsx')
    if not source.exists(): pytest.skip('trial.xlsx unavailable')
    records=parse_excel(source.read_bytes())
    client.get('/v3')  # Initializes V3 sidecar tables in the test DB.
    def confirm(selected, rows=records):
        token=uuid.uuid4().hex
        with db._conn() as c:
            c.execute('INSERT INTO v3_import_batches(token,shift_id,filename,created_at,rows_json) VALUES (?,?,?,?,?)',
                      (token,sid,'test.xlsx',datetime.now().isoformat(timespec='seconds'),json.dumps(rows)))
        response=client.post(f'/v3/import/{token}/confirm',data={'shift_id':str(sid),'techs':selected},follow_redirects=False)
        return response
    selected=['Ssaku, Isaac','Melvin, Douglas','Bawi, Mang','Popoola, Waheed']
    assert confirm(selected).status_code==303
    assert db.get_pm_count(sid)==8
    entry=next(e for e in db.get_entries(sid) if e['tech_name']=='Ssaku, Isaac')
    assert entry['excluded']==0
    assert confirm(selected[1:]).status_code==303
    assert db.get_entry(entry['id'])['excluded']==1
    assert db.get_pm_count(sid)==6
    assert confirm(selected).status_code==303
    assert db.get_entry(entry['id'])['excluded']==0
    assert db.get_pm_count(sid)==8
    # When a PM is reopened, no longer count it as a completed PM.
    changed=[{**r,'status':'Open'} if r['rfr']=='Preventive Maintenance' and r['tech_name']=='Ssaku, Isaac' else r for r in records]
    assert confirm(selected,changed).status_code==303
    assert db.get_pm_count(sid)==6


def test_loto_with_missing_keys_not_verified(app_client):
    client,sid=app_client
    result=client.post('/v3/loto',data={'shift_id':sid,'verified':'1','missing_keys':'Key #2'})
    assert result.status_code==200
    assert db.get_loto(sid)['loto_verified']==0
    assert db.get_loto(sid)['loto_missing_keys']=='Key #2'


def test_unfamiliar_rfr_imports_and_reopened_wo_is_excluded(app_client):
    """Full route/database path: new RFR included, then reopened excluded safely."""
    import json, uuid
    client,sid=app_client
    client.get('/v3')
    base={"wo_number":"98765","tech_name":"Ssaku, Isaac","asset_id":"ASSET987",
          "status":"Closed","rfr":"Safety Inspection","reported_issue":"Check guarding",
          "cause":"","remarks":"Guard inspected","start_at":"2026-10-20T18:00:00"}
    def send(rows):
        token=uuid.uuid4().hex
        with db._conn() as c:
            c.execute('INSERT INTO v3_import_batches(token,shift_id,filename,created_at,rows_json) VALUES (?,?,?,?,?)',
                      (token,sid,'sample.xlsx',datetime.now().isoformat(timespec='seconds'),json.dumps(rows)))
        return client.post(f'/v3/import/{token}/confirm',data={'shift_id':str(sid),'techs':['Ssaku, Isaac']},follow_redirects=True)
    first=send([base])
    assert first.status_code==200
    assert 'Safety Inspection' in first.text
    assert 'ASSET987' in first.text
    imported=next(e for e in db.get_entries(sid) if e['wo_number']=='98765')
    assert imported['excluded']==0
    second=send([{**base,'status':'Open'}])
    assert second.status_code==200
    assert db.get_entry(imported['id'])['excluded']==1
    third=send([base])
    assert third.status_code==200
    assert db.get_entry(imported['id'])['excluded']==0


def test_cumulative_reimport_keeps_work_orders_absent_from_later_file(app_client):
    """A new daily export must not silently remove older completed WOs or PMs."""
    import json, uuid
    client,sid=app_client
    client.get('/v3')
    base={"tech_name":"Tech A","asset_id":"M1","status":"Closed",
          "reported_issue":"Reported issue","cause":"WRONG COLUMN",
          "remarks":"Fixed","start_at":"2026-10-20T18:00:00"}
    def confirm(rows):
        token=uuid.uuid4().hex
        with db._conn() as c:
            c.execute('INSERT INTO v3_import_batches(token,shift_id,filename,created_at,rows_json) VALUES (?,?,?,?,?)',
                      (token,sid,'test.xlsx',datetime.now().isoformat(timespec='seconds'),json.dumps(rows)))
        response=client.post(f'/v3/import/{token}/confirm',data={'shift_id':sid,'techs':['Tech A']},follow_redirects=False)
        assert response.status_code==303,response.text
    confirm([{**base,'wo_number':'100','rfr':'Proactive Repair'},
             {**base,'wo_number':'200','rfr':'Preventive Maintenance'}])
    assert db.get_pm_count(sid)==1
    assert any(e['wo_number']=='100' and not e['excluded'] for e in db.get_entries(sid))
    confirm([{**base,'wo_number':'300','rfr':'Work Request'}])
    assert db.get_pm_count(sid)==1
    active={e['wo_number'] for e in db.get_entries(sid) if not e['excluded']}
    assert active=={'100','300'}


def test_s2_import_route_filters_out_s1_work(app_client):
    """S1 and S2 overlap, but the S2 import must not absorb S1 work."""
    import json, uuid
    client,sid=app_client
    client.get('/v3')
    rows=[
        {"wo_number":"S2-ONLY","asset_id":"M2","tech_name":"Tech S2","status":"Closed",
         "rfr":"Proactive Repair","reported_issue":"S2 repair","remarks":"RTS","cause":"",
         "start_at":"2026-10-20T18:00:00","in_service_at":"2026-10-20T18:10:00"},
        {"wo_number":"S1-ONLY","asset_id":"M1","tech_name":"Tech S1","status":"Closed",
         "rfr":"Proactive Repair","reported_issue":"S1 repair","remarks":"RTS","cause":"",
         "start_at":"2026-10-20T07:00:00","in_service_at":"2026-10-20T07:10:00"},
    ]
    token=uuid.uuid4().hex
    with db._conn() as c:
        c.execute("INSERT INTO v3_import_batches(token,shift_id,filename,created_at,rows_json) VALUES (?,?,?,?,?)",
                  (token,sid,"all_shifts.xlsx",datetime.now().isoformat(timespec="seconds"),json.dumps(rows)))
    response=client.post(f'/v3/import/{token}/confirm',
                         data={'shift_id':str(sid),'techs':['Tech S1','Tech S2']},
                         follow_redirects=False)
    assert response.status_code==303, response.text[:500]
    wos={e['wo_number'] for e in db.get_entries(sid) if not e['excluded']}
    assert "S2-ONLY" in wos
    assert "S1-ONLY" not in wos


def test_s3_import_route_keeps_friday_morning_ray_work(tmp_path):
    """Full route and DB: Thursday S3 owns Ray's Friday overlap work."""
    import json, uuid
    db.set_db_path(tmp_path/'s3_route.db')
    db.init_db()
    sid=db.start_shift('2026-10-22','S3','Test MST')
    app=FastAPI()
    app.mount('/static',StaticFiles(directory=str(DIR/'static')),name='static')
    templates=Jinja2Templates(directory=str(DIR/'templates'))
    app.include_router(create_v3_router(
        templates,
        is_lead_authed=lambda req: req.cookies.get('lead_auth')=='test',
        can_manage_shift=lambda req,candidate: candidate==sid,
        resolve_my_shift=lambda req: db.get_shift(sid),
        shift_window=lambda code,d:(datetime.combine(d,time(20,30)),
                                    datetime.combine(d+timedelta(days=1),time(8))),
        lead_redirect=lambda dest: RedirectResponse('/lead/login?next='+dest,status_code=303)))
    rows=[
        {"wo_number":"S3-EARLY","asset_id":"BALER0002","tech_name":"Stamps, Ray",
         "status":"Closed","rfr":"Preventive Maintenance","reported_issue":"PM",
         "remarks":"Completed","cause":"","start_at":"2026-10-23T03:00:00",
         "in_service_at":"2026-10-23T03:10:00"},
        {"wo_number":"28025279","asset_id":"VOIDFILL1201","tech_name":"Stamps, Ray",
         "status":"Closed","rfr":"Preventive Maintenance","reported_issue":"PM",
         "remarks":"Completed","cause":"","start_at":"2026-10-23T05:55:00",
         "in_service_at":"2026-10-23T06:53:00"},
        {"wo_number":"S4-ONLY","asset_id":"M4","tech_name":"Tech S4",
         "status":"Closed","rfr":"Proactive Repair","reported_issue":"S4 repair",
         "remarks":"RTS","cause":"","start_at":"2026-10-23T08:00:00",
         "in_service_at":"2026-10-23T08:10:00"},
    ]
    with TestClient(app) as client:
        client.cookies.set('lead_auth','test')
        client.get('/v3',params={'shift_id':sid})
        token=uuid.uuid4().hex
        with db._conn() as c:
            c.execute("INSERT INTO v3_import_batches(token,shift_id,filename,created_at,rows_json) VALUES (?,?,?,?,?)",
                      (token,sid,"mixed.xlsx",datetime.now().isoformat(timespec="seconds"),json.dumps(rows)))
        response=client.post(f'/v3/import/{token}/confirm',
                             data={'shift_id':str(sid),'techs':['Stamps, Ray','Tech S4']},
                             follow_redirects=False)
        assert response.status_code==303, response.text[:500]
        assert db.get_pm_count(sid)==2
        assert all(e['wo_number']!="S4-ONLY" for e in db.get_entries(sid))


def test_duplicate_completed_wo_cannot_credit_another_shift(app_client):
    """A previously credited WO cannot be counted as an S2 PM or happening."""
    import json,uuid
    client,sid=app_client
    client.get('/v3')
    other=db.start_shift('2026-10-20','S1','Other MST')
    db.upsert_pm_wo(other,'DUPLICATE-PM','M1','Tech S1')
    rows=[{"wo_number":"DUPLICATE-PM","asset_id":"M1","tech_name":"Tech S2",
           "status":"Closed","rfr":"Preventive Maintenance","reported_issue":"PM",
           "cause":"","remarks":"Done","start_at":"2026-10-20T18:00:00",
           "in_service_at":"2026-10-20T18:15:00"}]
    token=uuid.uuid4().hex
    with db._conn() as c:
        c.execute("INSERT INTO v3_import_batches(token,shift_id,filename,created_at,rows_json) VALUES (?,?,?,?,?)",
                  (token,sid,'dup.xlsx',datetime.now().isoformat(timespec='seconds'),json.dumps(rows)))
    response=client.post(f'/v3/import/{token}/confirm',
                         data={'shift_id':str(sid),'techs':['Tech S2']},
                         follow_redirects=False)
    assert response.status_code==303
    screen=client.get(f"/v3?shift_id={sid}")
    assert screen.status_code==200
    assert "Already credited to another shift (not counted)" in screen.text
    assert db.get_pm_count(sid)==0
    assert db.get_pm_count(other)==1
