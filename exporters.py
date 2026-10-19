"""DOCX and PDF exporters for shift passdown reports."""
from __future__ import annotations
import io
from datetime import datetime
from report_models import ReportData
from pm_table_style import pm_zebra_hex, PM_HEADER_HEX


def _fmt(iso: str) -> str:
    try:
        return datetime.strptime(iso, "%Y-%m-%d").strftime("%B %d, %Y")
    except Exception:
        return iso


# ── DOCX ───────────────────────────────────────────────────────────────────────────────

def export_docx(data: ReportData, pending_pms: list[dict] | None = None, loto_missing: str = "") -> bytes:
    """Build a branded DOCX matching the PDF layout and return as bytes."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml
    from lxml import etree

    # ── Walmart brand palette ────────────────────────────────────────────
    WM_BLUE  = RGBColor(0x00, 0x53, 0xE2)
    WM_LBLUE = RGBColor(0xBF, 0xDB, 0xFE)
    WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
    BLACK    = RGBColor(0x11, 0x18, 0x27)
    MED_GRAY = RGBColor(0x37, 0x41, 0x51)
    LGRAY    = RGBColor(0x6B, 0x72, 0x80)
    GREEN    = RGBColor(0x2A, 0x87, 0x03)
    RED      = RGBColor(0xEA, 0x11, 0x00)

    doc = Document()
    for sec in doc.sections:
        sec.left_margin   = Inches(0.85)
        sec.right_margin  = Inches(0.85)
        sec.top_margin    = Inches(0.5)
        sec.bottom_margin = Inches(0.75)

    # ── XML helpers ──────────────────────────────────────────────────────

    def _shade(cell, fill: str) -> None:
        shd = parse_xml(
            f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'
        )
        cell._tc.get_or_add_tcPr().append(shd)

    def _cell_pad(cell, top=60, bottom=60, left=100, right=100) -> None:
        cell._tc.get_or_add_tcPr().append(parse_xml(
            f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/></w:tcMar>'
        ))

    def _tbl_borders(tbl, outer: str | None = None) -> None:
        none = 'w:val="none" w:sz="0" w:space="0" w:color="auto"'
        box  = f'w:val="single" w:sz="12" w:space="0" w:color="{outer}"' if outer else none
        tbl._tbl.tblPr.append(parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top {box}/><w:left {box}/><w:bottom {box}/><w:right {box}/>'
            f'<w:insideH {none}/><w:insideV {none}/>'
            f'</w:tblBorders>'
        ))

    def _para_rule(para, color: str = "0053e2") -> None:
        para._element.get_or_add_pPr().append(parse_xml(
            f'<w:pBdr {nsdecls("w")}>'
            f'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="{color}"/>'
            f'</w:pBdr>'
        ))

    def _page_field(para) -> None:
        """Append a live PAGE number field to a paragraph."""
        def _fld(typ):
            r = etree.SubElement(para._element, qn("w:r"))
            fc = etree.SubElement(r, qn("w:fldChar"))
            fc.set(qn("w:fldCharType"), typ)
        _fld("begin")
        r = etree.SubElement(para._element, qn("w:r"))
        t = etree.SubElement(r, qn("w:instrText"))
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = " PAGE "
        _fld("separate")
        r = etree.SubElement(para._element, qn("w:r"))
        etree.SubElement(r, qn("w:t")).text = "1"
        _fld("end")

    # ── Banner (blue table, 2 rows) ──────────────────────────────────────
    dc = f"  |  {data.dc_number}" if data.dc_number else ""
    banner = doc.add_table(rows=2, cols=1)
    _tbl_borders(banner)

    top_cell = banner.rows[0].cells[0]
    _shade(top_cell, "0053e2")
    _cell_pad(top_cell, top=110, bottom=30, left=200, right=200)
    p = top_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"WALMART RELIABILITY OPERATIONS{dc}")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(15), WHITE

    sub_cell = banner.rows[1].cells[0]
    _shade(sub_cell, "0053e2")
    _cell_pad(sub_cell, top=20, bottom=110, left=200, right=200)
    # Yellow spark accent on bottom of banner
    sub_cell._tc.get_or_add_tcPr().append(parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:bottom w:val="single" w:sz="18" w:color="ffc220"/>'
        f'</w:tcBorders>'
    ))
    p = sub_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SHIFT PASSDOWN REPORT")
    r.font.size, r.font.color.rgb = Pt(9), WM_LBLUE

    doc.add_paragraph()  # spacer

    # ── Metadata table ───────────────────────────────────────────────────
    meta_data = [("DATE", _fmt(data.date)), ("SHIFT", data.shift_type), ("PREPARED BY", data.lead_name)]
    mt = doc.add_table(rows=3, cols=2)
    mt.columns[0].width = Inches(1.4)
    mt.columns[1].width = Inches(5.2)
    _tbl_borders(mt, outer="0053e2")
    for i, (lbl, val) in enumerate(meta_data):
        lc, vc = mt.rows[i].cells
        fill = "eff6ff" if i % 2 == 0 else "f8faff"
        for cell in (lc, vc):
            _shade(cell, fill)
            _cell_pad(cell)
        lr = lc.paragraphs[0].add_run(lbl)
        lr.bold, lr.font.size, lr.font.color.rgb = True, Pt(9), LGRAY
        vr = vc.paragraphs[0].add_run(val)
        vr.font.size, vr.font.color.rgb = Pt(10), BLACK

    # ── Section helpers ──────────────────────────────────────────────────

    def _sec(label: str) -> None:
        doc.add_paragraph()
        p = doc.add_paragraph()
        r = p.add_run(label.upper())
        r.bold, r.font.size, r.font.color.rgb = True, Pt(10), WM_BLUE
        _para_rule(p)
        p.paragraph_format.space_after = Pt(2)

    def _none(msg: str = "None this shift.") -> None:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run(msg)
        r.italic, r.font.size, r.font.color.rgb = True, Pt(10), LGRAY

    def _item(text: str, color: RGBColor = MED_GRAY, bold: bool = False) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = bold, Pt(10), color

    # ── Vendors ──────────────────────────────────────────────────────────
    _sec("Vendors")
    if data.vendors:
        for v in data.vendors: _item(f"\u25b6  {v}")
    else:
        _none("No vendor activity recorded during this shift.")

    # ── Shift Happenings ─────────────────────────────────────────────────
    _sec("Shift Happenings")
    if getattr(data, "shift_notes", "") and data.shift_notes.strip():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(data.shift_notes.strip())
        r.font.size, r.font.color.rgb = Pt(10), MED_GRAY
    if data.happenings:
        for h in data.happenings:
            _item(h.title, color=BLACK, bold=True)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(5)
            r = p.add_run(h.paragraph)
            r.font.size, r.font.color.rgb = Pt(10), MED_GRAY
    else:
        if not (getattr(data, "shift_notes", "") and data.shift_notes.strip()):
            _none("No shift happenings recorded.")

    # ── PM's Needing to be Completed ────────────────────────
    _sec("PM's Needing to be Completed")
    if pending_pms:
        from docx.enum.text import WD_ALIGN_PARAGRAPH as _WAP
        from docx.oxml.ns import nsdecls as _nsd
        from docx.oxml import parse_xml as _pxml
        pm_tbl = doc.add_table(rows=1, cols=5)
        pm_tbl.style = "Table Grid"
        hdr_cells = pm_tbl.rows[0].cells
        for ci, lbl in enumerate(["WO #", "Asset", "PM Due Date", "Promise Date/Time", "WO Date/Time"]):
            hdr_cells[ci].text = lbl
            r = hdr_cells[ci].paragraphs[0].runs[0]
            r.bold, r.font.size, r.font.color.rgb = True, Pt(9), WHITE
            _shade(hdr_cells[ci], PM_HEADER_HEX)
            _cell_pad(hdr_cells[ci], top=50, bottom=50)
        for i, r in enumerate(pending_pms):
            row = pm_tbl.add_row()
            vals = [r.get("wo_number",""), r.get("asset_num",""), r.get("pm_due_date",""),
                    r.get("promise_date",""), r.get("wo_date","")]
            bg = pm_zebra_hex(i)
            for ci, v in enumerate(vals):
                row.cells[ci].text = v
                run = row.cells[ci].paragraphs[0].runs[0]
                run.font.size = Pt(9)
                run.bold = ci in (1, 2)
                _shade(row.cells[ci], bg)
                _cell_pad(row.cells[ci], top=40, bottom=40)
        doc.add_paragraph()
    elif data.ro:
        for ro in data.ro: _item(f"\u25b6  {ro}")
    else:
        _none("No PMs recorded for this shift.")

    # ── Returned to Service ──────────────────────────────────────────────
    _sec("Returned to Service (RTS)")
    if data.rts:
        for u in data.rts:
            _item(f"\u2713  {u}  \u2014  Returned to Service", color=GREEN, bold=True)
    else:
        _none("No units returned to service this shift.")

    # ── Out of Service ───────────────────────────────────────────────────
    _sec("Current Out of Service Equipment")
    if data.oos:
        for o in data.oos:
            _item(f"[OOS]  {o.unit}", color=RED, bold=True)
    else:
        _none("No equipment currently out of service.")

    # ── Word footer (shift info left, page number right) ─────────────────
    fp = doc.sections[0].footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.3), WD_TAB_ALIGNMENT.RIGHT)
    lr = fp.add_run(f"{data.shift_type}  |  {_fmt(data.date)}  |  MST: {data.lead_name}")
    lr.font.size, lr.font.color.rgb = Pt(8), LGRAY
    fp.add_run("\t")
    pr = fp.add_run("Page ")
    pr.font.size, pr.font.color.rgb = Pt(8), LGRAY
    _page_field(fp)

    # ── LOTO + Signature ──────────────────────────────────────────────────────
    _sec("LOTO Keys")
    _item("All LOTO Keys Verified \u2713", color=GREEN, bold=True)
    if loto_missing:
        _item(f"Missing keys: {loto_missing}", color=RED, bold=True)

    # Signature
    doc.add_paragraph()
    p = doc.add_paragraph()
    _para_rule(p, color="e5e7eb")
    p.paragraph_format.space_after = Pt(4)
    sig_p = doc.add_paragraph()
    sig_p.paragraph_format.left_indent = Inches(0)
    r = sig_p.add_run(data.lead_name)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(12), BLACK
    title_p = doc.add_paragraph()
    r2 = title_p.add_run("MST Lead")
    r2.font.size, r2.font.color.rgb = Pt(10), MED_GRAY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF ────────────────────────────────────────────────────────────────────────────────

def export_pdf(data: ReportData, pending_pms: list[dict] | None = None, loto_missing: str = "") -> bytes:
    """Build a professionally branded PDF using ReportLab and return as bytes."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer,
        HRFlowable, Table, TableStyle, KeepTogether,
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    # ── Walmart brand palette ──────────────────────────────────────────────
    WM_BLUE    = colors.HexColor("#0053e2")
    WM_YELLOW  = colors.HexColor("#ffc220")
    BLACK      = colors.HexColor("#111827")
    MED_GRAY   = colors.HexColor("#374151")
    LGRAY      = colors.HexColor("#6b7280")
    BG_META    = colors.HexColor("#eff6ff")   # faint blue tint for meta box
    BG_ROW_ALT = colors.HexColor("#f8faff")
    GREEN      = colors.HexColor("#2a8703")
    RED        = colors.HexColor("#ea1100")
    RULE_COLOR = colors.HexColor("#bfdbfe")   # soft blue rule
    WHITE      = colors.white

    PAGE_W, PAGE_H = letter
    BANNER_H       = 1.1 * inch
    ACCENT_H       = 5                         # yellow strip height in pts
    L_MARGIN       = 0.75 * inch
    R_MARGIN       = 0.75 * inch
    CONTENT_W      = PAGE_W - L_MARGIN - R_MARGIN

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        leftMargin=L_MARGIN, rightMargin=R_MARGIN,
        topMargin=BANNER_H + 0.35 * inch,
        bottomMargin=0.75 * inch,
    )

    # ── Branded header + footer drawn on every page ────────────────────────
    def _draw_page(canvas, _doc):
        canvas.saveState()

        # Blue banner
        canvas.setFillColor(WM_BLUE)
        canvas.rect(0, PAGE_H - BANNER_H, PAGE_W, BANNER_H, fill=1, stroke=0)

        # Yellow accent strip at bottom of banner
        canvas.setFillColor(WM_YELLOW)
        canvas.rect(0, PAGE_H - BANNER_H, PAGE_W, ACCENT_H, fill=1, stroke=0)

        # Organisation name
        canvas.setFillColor(WHITE)
        canvas.setFont("Helvetica-Bold", 15)
        dc_suffix = f"  |  {data.dc_number}" if data.dc_number else ""
        canvas.drawCentredString(
            PAGE_W / 2, PAGE_H - 0.47 * inch,
            f"WALMART RELIABILITY OPERATIONS{dc_suffix}",
        )

        # Report sub-title
        canvas.setFont("Helvetica", 10)
        canvas.setFillColor(colors.HexColor("#bfdbfe"))  # soft blue
        canvas.drawCentredString(
            PAGE_W / 2, PAGE_H - 0.72 * inch,
            "SHIFT PASSDOWN REPORT",
        )

        # Footer rule
        footer_y = 0.55 * inch
        canvas.setStrokeColor(LGRAY)
        canvas.setLineWidth(0.5)
        canvas.line(L_MARGIN, footer_y, PAGE_W - R_MARGIN, footer_y)

        # Footer text
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(LGRAY)
        canvas.drawString(
            L_MARGIN, footer_y - 0.17 * inch,
            f"{data.shift_type}  |  {_fmt(data.date)}  |  MST: {data.lead_name}",
        )
        canvas.drawRightString(
            PAGE_W - R_MARGIN, footer_y - 0.17 * inch,
            f"Page {_doc.page}",
        )

        canvas.restoreState()

    # ── Paragraph styles ──────────────────────────────────────────────────
    base = getSampleStyleSheet()

    def S(name: str, **kw) -> ParagraphStyle:
        return ParagraphStyle(name, parent=base["Normal"], **kw)

    META_LBL  = S("ML", fontSize=9,    textColor=LGRAY,    fontName="Helvetica-Bold",    spaceAfter=0)
    META_VAL  = S("MV", fontSize=10.5, textColor=BLACK,    fontName="Helvetica",          spaceAfter=0)
    SECH      = S("SH", fontSize=10,   textColor=WM_BLUE,  fontName="Helvetica-Bold",
                  spaceBefore=16, spaceAfter=3, leading=14)
    BODY      = S("B",  fontSize=10.5, textColor=MED_GRAY, fontName="Helvetica",
                  leading=16.5, spaceAfter=4, leftIndent=16)
    ITEM_TITLE = S("IT", fontSize=10.5, textColor=BLACK,   fontName="Helvetica-Bold",
                   spaceBefore=6, spaceAfter=2, leftIndent=16)
    NONE_S    = S("NS", fontSize=10,   textColor=LGRAY,    fontName="Helvetica-Oblique",
                  leftIndent=16, spaceAfter=4)
    RTS_S     = S("RT", fontSize=10.5, textColor=GREEN,    fontName="Helvetica-Bold",
                  leftIndent=16, spaceAfter=4)
    OOS_S     = S("OO", fontSize=10.5, textColor=RED,      fontName="Helvetica-Bold",
                  leftIndent=16, spaceAfter=4)

    # ── Story ─────────────────────────────────────────────────────────────
    story: list = []

    # -- Metadata info box (alternating-row table) --------------------------
    meta_rows = [
        [Paragraph("DATE", META_LBL),        Paragraph(_fmt(data.date), META_VAL)],
        [Paragraph("SHIFT", META_LBL),       Paragraph(data.shift_type, META_VAL)],
        [Paragraph("PREPARED BY", META_LBL), Paragraph(data.lead_name, META_VAL)],
    ]
    col_w = [1.25 * inch, CONTENT_W - 1.25 * inch]
    meta_tbl = Table(meta_rows, colWidths=col_w)
    meta_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, -1), BG_META),
        ("BACKGROUND",    (0, 1), (-1,  1), BG_ROW_ALT),
        ("TOPPADDING",    (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LEFTPADDING",   (0, 0), (-1, -1), 10),
        ("RIGHTPADDING",  (0, 0), (-1, -1), 10),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW",     (0, 0), (-1, -2), 0.5, RULE_COLOR),
        ("BOX",           (0, 0), (-1, -1), 0.75, WM_BLUE),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 0.18 * inch))

    # -- Section helpers ----------------------------------------------------
    def add_section(label: str) -> None:
        story.append(Spacer(1, 0.04 * inch))
        story.append(Paragraph(label.upper(), SECH))
        story.append(HRFlowable(
            width="100%", thickness=1.5,
            color=WM_BLUE, spaceAfter=5,
        ))

    def add_none(msg: str = "None this shift.") -> None:
        story.append(Paragraph(msg, NONE_S))

    # ── VENDORS ───────────────────────────────────────────────────────────
    add_section("Vendors")
    if data.vendors:
        for v in data.vendors:
            story.append(Paragraph(f"&#9658;&#160;&#160;{v}", BODY))
    else:
        add_none("No vendor activity recorded during this shift.")

    # ── SHIFT HAPPENINGS ──────────────────────────────────────────────────
    add_section("Shift Happenings")
    if getattr(data, "shift_notes", "") and data.shift_notes.strip():
        story.append(Paragraph(data.shift_notes.strip(), BODY))
    if data.happenings:
        for h in data.happenings:
            story.append(KeepTogether([
                Paragraph(h.title, ITEM_TITLE),
                Paragraph(h.paragraph, BODY),
                Spacer(1, 2),
            ]))
    else:
        if not (getattr(data, "shift_notes", "") and data.shift_notes.strip()):
            add_none("No shift happenings recorded.")

    # ── PM'S NEEDING TO BE COMPLETED ────────────────────────
    add_section("PM's Needing to be Completed")
    if pending_pms:
        PM_HDR = S("PH", fontSize=9,  textColor=WHITE,     fontName="Helvetica-Bold",  spaceAfter=0)
        PM_TD  = S("PT", fontSize=9,  textColor=MED_GRAY,  fontName="Helvetica",        spaceAfter=0)
        PM_TDB = S("PTB",fontSize=9,  textColor=BLACK,     fontName="Helvetica-Bold",   spaceAfter=0)
        headers = [[Paragraph(h, PM_HDR) for h in ["WO #", "Asset", "PM Due Date", "Promise Date/Time", "WO Date/Time"]]]
        rows = []
        for i, r in enumerate(pending_pms):
            row = [
                Paragraph(r.get("wo_number",""),  PM_TD),
                Paragraph(r.get("asset_num",""),   PM_TDB),
                Paragraph(r.get("pm_due_date",""), PM_TDB),
                Paragraph(r.get("promise_date",""),PM_TD),
                Paragraph(r.get("wo_date",""),     PM_TD),
            ]
            rows.append(row)
        col_w = [1.0*inch, 1.1*inch, 1.1*inch, 1.6*inch, 1.4*inch]
        pm_tbl = Table(headers + rows, colWidths=col_w)
        style_cmds = [
            ("BACKGROUND",    (0,0),(-1,0),  WM_BLUE),
            ("TEXTCOLOR",     (0,0),(-1,0),  WHITE),
            ("TOPPADDING",    (0,0),(-1,-1), 5),
            ("BOTTOMPADDING", (0,0),(-1,-1), 5),
            ("LEFTPADDING",   (0,0),(-1,-1), 6),
            ("RIGHTPADDING",  (0,0),(-1,-1), 6),
            ("GRID",          (0,0),(-1,-1), 0.5, colors.HexColor("#d1d5db")),
            ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ]
        for i, r in enumerate(pending_pms):
            style_cmds.append(("BACKGROUND",(0,i+1),(-1,i+1),colors.HexColor(f"#{pm_zebra_hex(i)}")))
        pm_tbl.setStyle(TableStyle(style_cmds))
        story.append(pm_tbl)
    elif data.ro:
        for r in data.ro:
            story.append(Paragraph(f"&#9658;&#160;&#160;{r}", BODY))
    else:
        add_none("No PMs recorded for this shift.")

    # ── RETURNED TO SERVICE (RTS) ─────────────────────────────────────────
    add_section("Returned to Service (RTS)")
    if data.rts:
        for u in data.rts:
            story.append(Paragraph(
                f"<font color='#2a8703'>&#10004;</font>&#160;&#160;"
                f"{u} &#8212; Returned to Service",
                RTS_S,
            ))
    else:
        add_none("No units returned to service this shift.")

    # ── CURRENT OUT OF SERVICE ────────────────────────────────────────────
    add_section("Current Out of Service Equipment")
    if data.oos:
        for o in data.oos:
            story.append(Paragraph(
                f"<font color='#ea1100'>&#9940;</font>&#160;&#160;{o.unit}",
                OOS_S,
            ))
    else:
        add_none("No equipment currently out of service.")

    story.append(Spacer(1, 0.25 * inch))

    # ── LOTO + Signature ─────────────────────────────────────────────────
    add_section("LOTO Keys")
    story.append(Paragraph(
        "<font color='#2a8703'>&#10004;</font>&#160;&#160;All LOTO Keys Verified",
        RTS_S,
    ))
    if loto_missing:
        story.append(Paragraph(
            f"<font color='#ea1100'>&#9888;</font>&#160;&#160;Missing: {loto_missing}",
            OOS_S,
        ))

    # Signature line
    story.append(Spacer(1, 0.3 * inch))
    story.append(HRFlowable(width="100%", thickness=0.5, color=LGRAY))
    story.append(Spacer(1, 0.08 * inch))
    SIG_NAME  = S("SN", fontSize=12, textColor=BLACK,    fontName="Helvetica-Bold")
    SIG_TITLE = S("ST", fontSize=10, textColor=MED_GRAY, fontName="Helvetica")
    story.append(Paragraph(data.lead_name, SIG_NAME))
    story.append(Paragraph("MST Lead",     SIG_TITLE))
    story.append(Spacer(1, 0.25 * inch))

    doc.build(story, onFirstPage=_draw_page, onLaterPages=_draw_page)
    return buf.getvalue()
